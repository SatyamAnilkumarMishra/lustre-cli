#!/usr/bin/env python3
"""Generate lustre-cli project Word report with code blocks and test placeholders."""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
OUTPUT = DOCS / "LUSTRE_CLI_PROJECT_REPORT.docx"
IMAGES = DOCS / "report_images"


def _ensure_docx():
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt, RGBColor
        return Document, WD_ALIGN_PARAGRAPH, Inches, Pt, RGBColor
    except ImportError:
        import subprocess
        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "-q"])
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt, RGBColor
        return Document, WD_ALIGN_PARAGRAPH, Inches, Pt, RGBColor


def add_title(doc, WD_ALIGN_PARAGRAPH, Pt):
    t = doc.add_heading("lustre-cli", 0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("iSCSI + Lustre Storage Management CLI")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.size = Pt(14)
    doc.add_paragraph("Complete Project Documentation")
    doc.add_paragraph("Platform: Ubuntu 22.04 LTS Server (64-bit)")
    doc.add_paragraph("Author: _________________________")
    doc.add_paragraph("Organization: _________________________")
    doc.add_paragraph("Date: _________________________")
    doc.add_page_break()


def add_code(doc, code: str, caption: str = ""):
    from docx.shared import Pt, Inches
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    if caption:
        p = doc.add_paragraph()
        r = p.add_run(caption)
        r.bold = True
        r.font.size = Pt(11)

    p = doc.add_paragraph()
    run = p.add_run(code.rstrip())
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    p.paragraph_format.left_indent = Inches(0.2)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F2F2")
    p._p.get_or_add_pPr().append(shading)
    doc.add_paragraph()


def add_image_placeholder(doc, Inches, title: str, hint: str):
    doc.add_paragraph(title).runs[0].bold = True
    p = doc.add_paragraph()
    p.add_run(f"[INSERT SCREENSHOT: {hint}]").italic = True
    doc.add_paragraph(
        "Tip: On Ubuntu VM, press Print Screen or use "
        "gnome-screenshot / scrot. Save PNG to docs/report_images/ "
        "and replace this placeholder in Word (Insert → Pictures)."
    )
    doc.add_paragraph()


def build():
    Document, WD_ALIGN_PARAGRAPH, Inches, Pt, RGBColor = _ensure_docx()
    IMAGES.mkdir(parents=True, exist_ok=True)
    doc = Document()

    # Styles
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    add_title(doc, WD_ALIGN_PARAGRAPH, Pt)

    # 1. Introduction
    doc.add_heading("1. Introduction", 1)
    doc.add_paragraph(
        "lustre-cli is a Linux command-line tool that integrates iSCSI networked block storage "
        "with the Lustre distributed parallel filesystem. It allows administrators to export block "
        "devices from a storage server, import them on a client via iSCSI, format and mount Lustre "
        "(MGS, MDT, OST roles), validate filesystem operation, benchmark performance, simulate "
        "failures, and tear down the environment cleanly."
    )
    doc.add_paragraph("Key technologies:")
    for item in [
        "iSCSI Target: targetcli (LIO)",
        "iSCSI Initiator: open-iscsi (iscsiadm)",
        "Lustre: mkfs.lustre, mount.lustre, lctl, lfs, lnetctl",
        "Benchmarking: fio, dd",
        "Language: Python 3.9+",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("1.1 Objectives", 2)
    for obj in [
        "Export block devices using iSCSI target setup",
        "Import devices on client using iSCSI initiator",
        "Deploy Lustre filesystem with MGS, MDT, and OST components",
        "Validate I/O, striping, and data integrity",
        "Benchmark sequential/random performance",
        "Simulate OST failure, bad configuration, and network issues",
        "Provide teardown and hard reset for clean lab environments",
    ]:
        doc.add_paragraph(obj, style="List Bullet")

    doc.add_page_break()

    # 2. Architecture
    doc.add_heading("2. System Architecture", 1)
    doc.add_paragraph(
        "The architecture has two layers. iSCSI presents remote block devices as local /dev/sdX "
        "devices. Lustre formats those devices as management (MGS), metadata (MDT), and object "
        "storage (OST) targets, then exposes a unified filesystem namespace to applications."
    )
    add_code(
        doc,
        """STORAGE SERVER                         CLIENT NODE
/dev/sdb, /dev/sdc  ──►  targetcli  ──►  iSCSI :3260
                              │                    │
                              └──── TCP/IP ─────────► iscsiadm → /dev/sdX
                                                         │
                                                         ▼
                                              mkfs.lustre (MGS/MDT/OST)
                                                         │
                                                         ▼
                                              /mnt/lustre/client""",
        "Figure 2.1 — Architecture flow (text diagram)",
    )

    doc.add_heading("2.1 Lustre Components", 2)
    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    rows = [
        ("Component", "Role"),
        ("MGS", "Management Server — stores filesystem configuration"),
        ("MDT", "Metadata Target — filenames, directories, permissions"),
        ("OST", "Object Storage Target — actual file data chunks"),
    ]
    for i, (a, b) in enumerate(rows):
        table.rows[i].cells[0].text = a
        table.rows[i].cells[1].text = b

    doc.add_page_break()

    # 3. Project structure
    doc.add_heading("3. Project Structure", 1)
    add_code(
        doc,
        """lustre-cli/
├── lustre_cli/           # Python package
│   ├── main.py           # CLI entry point
│   ├── target.py         # Module 1: iSCSI target
│   ├── initiator.py      # Module 2: iSCSI initiator
│   ├── deploy.py         # Module 3: Lustre deployment
│   ├── validate.py       # Module 4: Validation
│   ├── benchmark.py      # Module 5: Benchmarking
│   ├── fault.py          # Module 6: Fault simulation
│   ├── teardown.py       # Module 7: Cleanup
│   ├── config.py         # YAML configuration
│   ├── utils.py          # Subprocess helpers
│   ├── deps.py           # Dependency checks
│   └── logging_util.py   # File logging
├── scripts/install.sh    # Ubuntu/RHEL installer
├── config/config.yaml.example
└── docs/                 # Documentation""",
        "Listing 3.1 — Directory layout",
    )

    doc.add_page_break()

    # 4. Installation Ubuntu
    doc.add_heading("4. Installation on Ubuntu 22.04 LTS", 1)
    doc.add_paragraph(
        "Download Ubuntu 22.04.5 LTS Server ISO from ubuntu.com/download/server. "
        "Install in VirtualBox/VMware with 4 GB RAM and 40 GB disk."
    )

    add_code(
        doc,
        """# Clone project
git clone https://github.com/YOUR_USERNAME/lustre-cli.git
cd lustre-cli
chmod +x scripts/install.sh
sudo ./scripts/install.sh

# Verify
lustre-cli --version
sudo lustre-cli check-deps""",
        "Listing 4.1 — Install commands on Ubuntu",
    )
    doc.add_paragraph(
        "Explanation: install.sh detects apt, installs targetcli-fb, open-iscsi, fio, Python "
        "dependencies, runs pip install for lustre-cli, creates /etc/lustre-cli/config.yaml "
        "and /var/log/lustre-cli.log."
    )

    add_image_placeholder(
        doc, Inches,
        "Figure 4.1 — Ubuntu terminal: lustre-cli check-deps output",
        "Run 'sudo lustre-cli check-deps' on Ubuntu and capture terminal showing all tools OK",
    )

    doc.add_page_break()

    # 5. Modules
    modules = [
        (
            "5. Module 1 — iSCSI Target (target.py)",
            "Creates iSCSI targets on the storage server using targetcli. Maps a block device "
            "to a LUN behind an IQN, binds a portal on IP:3260, disables authentication for lab use, "
            "and saves configuration for reboot persistence.",
            """def cmd_create(device, lun, portal_ip=None, portal_port=None, backstore_type="block"):
    require_root()
    check_tools("target")
    iqn = _iqn_for_lun(cfg, lun)   # e.g. iqn.2024-05.com.lustre-cli:lun0
    cmds = [
        f"/backstores/block create name=bs_lun{lun} {device}",
        f"/iscsi create {iqn}",
        f"/iscsi/{iqn}/tpg1/luns create /backstores/block/bs_lun{lun}",
        f"/iscsi/{iqn}/tpg1/portals create",
        "saveconfig",
    ]
    _targetcli_batch(cmds)""",
            """# Create target on /dev/sdb as LUN 0
sudo lustre-cli target create -d /dev/sdb --lun 0

# List targets
sudo lustre-cli target list

# Delete target
sudo lustre-cli target delete --lun 0""",
            "Figure 5.1 — target create and target list on Ubuntu",
            "Screenshot of 'sudo lustre-cli target list' showing IQN and device mapping",
        ),
        (
            "6. Module 2 — iSCSI Initiator (initiator.py)",
            "Discovers targets via iscsiadm, logs in to a specific IQN, sets node.startup=automatic "
            "for reboot persistence, detects the attached /dev/sdX device, and saves session info to config.",
            """def cmd_login(host, iqn, port=3260):
    run_cmd(["iscsiadm", "-m", "node", "-T", iqn, "-p", portal, "--login"])
    run_cmd(["iscsiadm", "-m", "node", "-T", iqn, "-p", portal,
             "-o", "update", "-n", "node.startup", "-v", "automatic"])
    device = _find_session_device(iqn)  # parses session -P 3 or lsblk TRAN=iscsi""",
            """sudo lustre-cli initiator discover --host 10.0.0.10
sudo lustre-cli initiator login --host 10.0.0.10 \\
  --iqn iqn.2024-05.com.lustre-cli:lun0
sudo lustre-cli initiator status
lsblk   # verify new iSCSI disk""",
            "Figure 6.1 — initiator discover and login",
            "Screenshot showing discover output and lsblk with iscsi transport",
        ),
        (
            "7. Module 3 — Lustre Deployment (deploy.py)",
            "Loads Lustre kernel modules (libcfs, lnet, lustre), configures LNet via lnetctl, "
            "runs mkfs.lustre for MGS/MDT/OST with correct flags, mounts each component, "
            "and mounts the client namespace at /mnt/lustre/client.",
            """# MGS format
mkfs.lustre --mgs --fsname=lustrefs --reformat /dev/sdb

# MDT format
mkfs.lustre --mdt --mgsnode=10.0.0.10@tcp --fsname=lustrefs --index=0 /dev/sdc

# OST format
mkfs.lustre --ost --mgsnode=10.0.0.10@tcp --fsname=lustrefs --index=0 /dev/sdd""",
            """sudo lustre-cli deploy format \\
  --mgs-device /dev/sdb --mdt-device /dev/sdc \\
  --ost-device /dev/sdd --mgsnode 10.0.0.10@tcp --force

sudo lustre-cli deploy mount
sudo lustre-cli deploy status""",
            "Figure 7.1 — deploy status showing Lustre mounts",
            "Screenshot of 'sudo lustre-cli deploy status' with findmnt and lctl dl output",
        ),
        (
            "8. Module 4 — Validation (validate.py)",
            "Three validation tests: basic file create/read/write, stripe configuration via "
            "lfs setstripe/getstripe, and SHA256 integrity check with OST distribution via lfs df.",
            """def cmd_integrity():
    data = os.urandom(1024 * 256)
    digest = hashlib.sha256(data).hexdigest()
    # write file on Lustre mount, read back, compare digest
    run_cmd(["lctl", "dl"])
    run_cmd(["lfs", "df", "-h"])""",
            """sudo lustre-cli validate basic
sudo lustre-cli validate stripe
sudo lustre-cli validate integrity""",
            "Figure 8.1 — validate basic PASS output",
            "Screenshot showing 'PASS: created ... wrote and verified'",
        ),
        (
            "9. Module 5 — Benchmarking (benchmark.py)",
            "Runs fio (or dd fallback) with varying stripe counts. Measures throughput (MB/s), "
            "IOPS, and latency. Saves JSON report to /var/lib/lustre-cli/benchmarks/.",
            """def cmd_run(runtime=None, stripe_counts=None, use_dd_fallback=False):
    stripes = stripe_counts or [1, ost_count]
    results = _run_fio_benchmark(mp, runtime_sec, stripes)
    report_path = out_dir / f"benchmark_{timestamp}.json"
    _print_table(results)""",
            """sudo lustre-cli benchmark run --runtime 30
sudo lustre-cli benchmark report""",
            "Figure 9.1 — Benchmark results table",
            "Screenshot of benchmark run output showing MB/s and IOPS table",
        ),
        (
            "10. Module 6 — Fault Simulation (fault.py)",
            "Simulates real failure scenarios: OST unmount, wrong MGS IP mount attempt, "
            "re-format existing device, wrong IQN login, unreachable portal discovery. "
            "Each fault logs timestamp and remediation steps.",
            """REMEDIATION = {
    "ost_offline": "Restore OST: verify iSCSI session, remount OST...",
    "bad_mgs": "Verify lustre.mgsnode IP and LNet connectivity...",
    "login_fail": "Check target IQN, firewall (tcp/3260)...",
}""",
            """sudo lustre-cli fault simulate-ost-failure --index 0
sudo lustre-cli fault simulate-bad-config
sudo lustre-cli fault simulate-network-drop""",
            "Figure 10.1 — Fault simulation with remediation message",
            "Screenshot showing FAULT log line and Remediation text",
        ),
        (
            "11. Module 7 — Teardown (teardown.py)",
            "Unmounts Lustre in correct order (client → OST → MDT → MGS), logs out iSCSI "
            "sessions, optionally wipefs devices. Hard reset also deletes targetcli config.",
            """def cmd_teardown(wipe=False):
    deploy.cmd_unmount()          # Lustre unmount
    initiator.cmd_logout(...)     # iSCSI logout
    if wipe: _wipe_devices(cfg)   # wipefs -a

def cmd_reset_hard():
    cmd_teardown(wipe=True)
    target.cmd_delete(...)      # remove all targets
    targetcli clearconfig""",
            """sudo lustre-cli teardown
sudo lustre-cli reset --hard""",
            "Figure 11.1 — Teardown complete message",
            "Screenshot showing teardown and clean lsblk output",
        ),
    ]

    for title, explanation, code_py, code_cli, fig_title, fig_hint in modules:
        doc.add_heading(title, 1)
        doc.add_paragraph(explanation)
        add_code(doc, code_py, "Key implementation:")
        add_code(doc, code_cli, "CLI usage on Ubuntu:")
        add_image_placeholder(doc, Inches, fig_title, fig_hint)
        doc.add_page_break()

    # 12. Core utilities
    doc.add_heading("12. Core Utilities", 1)

    doc.add_heading("12.1 Configuration (config.py)", 2)
    add_code(
        doc,
        """DEFAULT_CONFIG_PATH = Path("/etc/lustre-cli/config.yaml")

def load_config(path=None):
    data = deepcopy(DEFAULTS)
    if cfg_path.is_file():
        loaded = yaml.safe_load(fh)
        _deep_merge(data, loaded)
    return data""",
        "Loads YAML config with defaults for iSCSI, Lustre, benchmark, and logging.",
    )

    doc.add_heading("12.2 Logging (logging_util.py)", 2)
    doc.add_paragraph(
        "All operations log to /var/log/lustre-cli.log with timestamps. "
        "Format: YYYY-MM-DD HH:MM:SS [LEVEL] lustre-cli: message"
    )

    doc.add_heading("12.3 CLI Entry Point (main.py)", 2)
    add_code(
        doc,
        """def main(argv=None):
    parser = build_parser()   # argparse with all subcommands
    args = parser.parse_args(argv)
    return _dispatch(args)    # routes to target/initiator/deploy/...""",
        "Uses argparse for --help on every subcommand. Returns non-zero exit code on CLIError.",
    )

    doc.add_page_break()

    # 13. Full Ubuntu test procedure
    doc.add_heading("13. Ubuntu Test Procedure (Step-by-Step)", 1)
    steps = [
        ("Step 1", "Install Ubuntu 22.04 LTS Server in VM"),
        ("Step 2", "Clone repo and run sudo ./scripts/install.sh"),
        ("Step 3", "Run sudo lustre-cli check-deps — all tools must pass"),
        ("Step 4", "Create loop devices OR use real disks for iSCSI"),
        ("Step 5", "sudo lustre-cli target create -d /dev/loop0 --lun 0"),
        ("Step 6", "sudo lustre-cli initiator discover --host 127.0.0.1"),
        ("Step 7", "sudo lustre-cli initiator login --host 127.0.0.1 --iqn <IQN>"),
        ("Step 8", "sudo lustre-cli deploy format ... && deploy mount"),
        ("Step 9", "Run validate basic, stripe, integrity"),
        ("Step 10", "Run benchmark run and benchmark report"),
        ("Step 11", "Run fault simulate-* commands"),
        ("Step 12", "Run teardown and document results"),
    ]
    for step, desc in steps:
        p = doc.add_paragraph()
        p.add_run(f"{step}: ").bold = True
        p.add_run(desc)

    add_code(
        doc,
        """# Single-machine lab — create test disks
sudo truncate -s 2G /var/lib/lustre-cli/disk0.img
sudo losetup -f /var/lib/lustre-cli/disk0.img
lsblk

export IP=$(hostname -I | awk '{print $1}')
sudo lustre-cli target create -d /dev/loop0 --lun 0
sudo lustre-cli initiator discover --host 127.0.0.1
sudo lustre-cli initiator login --host 127.0.0.1 \\
  --iqn iqn.2024-05.com.lustre-cli:lun0""",
        "Listing 13.1 — Single-node Ubuntu lab commands",
    )

    doc.add_page_break()

    # 14. Test results table
    doc.add_heading("14. Test Results (Fill After Ubuntu Testing)", 1)
    doc.add_paragraph(
        "Run tests on Ubuntu and fill this table. Insert screenshots in Section 4–11 placeholders."
    )
    t = doc.add_table(rows=10, cols=4)
    t.style = "Table Grid"
    headers = ["Test", "Command", "Result (PASS/FAIL)", "Notes"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].text = h
    tests = [
        ("Dependency check", "lustre-cli check-deps", "", ""),
        ("Target create", "target create -d /dev/loop0 --lun 0", "", ""),
        ("Initiator login", "initiator login ...", "", ""),
        ("Lustre format", "deploy format ...", "", ""),
        ("Lustre mount", "deploy mount", "", ""),
        ("Validation", "validate basic", "", ""),
        ("Benchmark", "benchmark run", "", ""),
        ("Fault test", "fault simulate-bad-config", "", ""),
        ("Teardown", "teardown", "", ""),
    ]
    for i, row in enumerate(tests, 1):
        for j, val in enumerate(row):
            t.rows[i].cells[j].text = val

    doc.add_page_break()

    # 15. Screenshot checklist
    doc.add_heading("15. Screenshot Checklist for Word Report", 1)
    shots = [
        "Ubuntu VM login / uname -a showing 22.04",
        "git clone and install.sh success",
        "lustre-cli check-deps all OK",
        "target list with IQN",
        "initiator status with session",
        "lsblk showing iscsi devices",
        "deploy status with lustre mounts",
        "validate basic PASS",
        "benchmark report table",
        "fault simulation remediation output",
        "tail /var/log/lustre-cli.log",
    ]
    for s in shots:
        doc.add_paragraph(s, style="List Number")

    doc.add_paragraph(
        "Save screenshots as PNG in docs/report_images/ then in Word: "
        "Insert → Pictures → select file → replace placeholder paragraphs."
    )

    doc.add_page_break()

    # 16. Conclusion
    doc.add_heading("16. Conclusion", 1)
    doc.add_paragraph(
        "lustre-cli successfully integrates iSCSI block storage export/import with Lustre "
        "filesystem deployment in a single command-line tool. All seven functional modules are "
        "implemented with logging, configuration persistence, dependency validation, and root "
        "privilege checks. The tool is designed for Ubuntu/RHEL Linux environments and supports "
        "lab testing, validation, benchmarking, and fault simulation for storage administration training."
    )

    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")
    return OUTPUT


if __name__ == "__main__":
    build()
